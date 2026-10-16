from __future__ import annotations

import io
from datetime import timedelta
from decimal import Decimal

import sqlalchemy as sa
from openpyxl import load_workbook

from app.domain.enums import DocumentType, Role, StaffType
from tests.conftest import Harness
from tests.fakes import invoice, make_jpeg


async def _upload_and_confirm(harness: Harness, headers: dict[str, str], seed: int, **data: object) -> str:
    harness.extractor.respond = lambda _r, _n: [invoice(invoice_number=f"F-{seed}", **data)]
    upload = await harness.client.post(
        "/api/v1/expenses", headers=headers, files={"file": ("f.jpg", make_jpeg(seed=seed), "image/jpeg")}
    )
    await harness.drain()
    expense_id = upload.json()["id"]
    confirmed = await harness.client.post(f"/api/v1/expenses/{expense_id}/confirm", headers=headers, json={"justification": "ok"})
    assert confirmed.status_code == 200, confirmed.text
    return expense_id


async def _scenario(harness: Harness):  # type: ignore[no-untyped-def]
    company = await harness.add_company()
    other = await harness.add_company("Forsa Panamá", "PA", "USD")
    accountant = await harness.add_employee(company, "conta@forsa.com", Role.ACCOUNTANT)
    approver = await harness.add_employee(company, "ana@forsa.com", Role.APPROVER)
    employee = await harness.add_employee(company, "juan@forsa.com", approver=approver, staff_type=StaffType.OPERATIONAL)
    acc_h, appr_h, emp_h = await harness.login(accountant), await harness.login(approver), await harness.login(employee)
    today = harness.container.clock.now().date()
    created = await harness.client.post(
        "/api/v1/advances",
        headers=emp_h,
        json={
            "company_id": str(company.id),
            "kind": "cash",
            "currency": "COP",
            "amount": "400000",
            "purpose": "Obra Medellín",
            "destination": "Medellín",
            "start_date": str(today - timedelta(days=45)),
            "end_date": str(today - timedelta(days=40)),
        },
    )
    advance_id = created.json()["id"]
    await harness.client.post(f"/api/v1/advances/{advance_id}/approve", headers=appr_h)
    return company, other, accountant, acc_h, emp_h, advance_id, today


async def test_pending_balances_with_age_and_spend(harness: Harness) -> None:
    _, other, _, acc_h, emp_h, advance_id, today = await _scenario(harness)
    day = today - timedelta(days=42)
    await _upload_and_confirm(harness, emp_h, 1, total=Decimal("45000"), invoice_date=day, category="meals")
    await _upload_and_confirm(harness, emp_h, 2, total=Decimal("120000"), invoice_date=day, category="lodging")

    balances = (await harness.client.get("/api/v1/reports/pending-balances", headers=acc_h)).json()["data"]
    assert len(balances) == 1 and balances[0]["advance_id"] == advance_id
    assert Decimal(balances[0]["pending"]) == Decimal("235000") and balances[0]["days"] == 40 and balances[0]["bucket"] == "+30"

    spend = await harness.client.get(
        "/api/v1/reports/spend", headers=acc_h, params={"start": str(day - timedelta(days=1)), "end": str(today)}
    )
    totals = {row["category"]: Decimal(row["total"]) for row in spend.json()["data"]}
    assert totals == {"lodging": Decimal("120000"), "meals": Decimal("45000")}

    excel = await harness.client.get("/api/v1/reports/pending-balances", headers=acc_h, params={"format": "xlsx"})
    sheet = load_workbook(io.BytesIO(excel.content)).active
    assert any("Juan Pérez" in [c for c in row] for row in sheet.iter_rows(values_only=True))

    # Permisos: el empleado no ve tableros y contabilidad no ve otras compañías.
    assert (await harness.client.get("/api/v1/reports/pending-balances", headers=emp_h)).status_code == 403
    foreign = await harness.client.get("/api/v1/reports/pending-balances", headers=acc_h, params={"company_id": str(other.id)})
    assert foreign.status_code == 404


async def test_budget_compliance_and_intercompany_invoices(harness: Harness) -> None:
    company, other, _, acc_h, emp_h, _, today = await _scenario(harness)
    admin = await harness.add_employee(company, "admin@forsa.com", Role.ADMIN)
    admin_h = await harness.login(admin)
    await harness.client.patch(f"/api/v1/companies/{company.id}", headers=admin_h, json={"tax_id": "900123456-7"})
    await harness.client.patch(f"/api/v1/companies/{other.id}", headers=admin_h, json={"tax_id": "155612345-2-2026"})
    zone = await harness.client.post(
        f"/api/v1/companies/{company.id}/travel-zones", headers=acc_h, json={"name": "Principales", "places": ["Cali"]}
    )
    await harness.client.post(
        f"/api/v1/companies/{company.id}/travel-rates",
        headers=acc_h,
        json={
            "zone_id": zone.json()["id"],
            "staff_type": "operational",
            "category_code": "meals",
            "amount": "50000",
            "currency": "COP",
            "unit": "per_day",
            "valid_from": "2026-01-01",
        },
    )
    trip = await harness.client.post(
        "/api/v1/advances",
        headers=emp_h,
        json={
            "company_id": str(company.id),
            "kind": "cash",
            "currency": "COP",
            "purpose": "Cali",
            "destination": "Cali",
            "start_date": str(today),
            "end_date": str(today + timedelta(days=1)),
        },
    )
    approver_h = await harness.login(await harness.add_employee(company, "jefe@forsa.com", Role.ACCOUNTANT))
    await harness.client.post(f"/api/v1/advances/{trip.json()['id']}/approve", headers=approver_h)
    expense = await _upload_and_confirm(
        harness,
        emp_h,
        3,
        total=Decimal("70000"),
        invoice_date=today,
        category="meals",
        document_type=DocumentType.ELECTRONIC_INVOICE,
        buyer_tax_id="155612345-2-2026",
    )
    # Con dos anticipos abiertos el gasto no se asigna solo: se pasa al de Cali.
    await harness.client.post(
        f"/api/v1/expenses/{expense}/confirm", headers=emp_h, json={"advance_id": trip.json()["id"], "justification": "ok"}
    )

    window = {"start": str(today - timedelta(days=1)), "end": str(today + timedelta(days=1))}
    budget = (await harness.client.get("/api/v1/reports/budget-compliance", headers=acc_h, params=window)).json()["data"]
    assert budget[0]["category"] == "meals" and Decimal(budget[0]["budget"]) == Decimal("100000")
    intercompany = (await harness.client.get("/api/v1/reports/intercompany-invoices", headers=acc_h, params=window)).json()
    assert intercompany["data"][0]["billed_to"] == "Forsa Panamá"


async def test_usage_and_accuracy(harness: Harness) -> None:
    company, _, _, acc_h, emp_h, _, today = await _scenario(harness)
    expense = await _upload_and_confirm(harness, emp_h, 4, total=Decimal("45000"), invoice_date=today - timedelta(days=41))
    await harness.client.patch(f"/api/v1/expenses/{expense}", headers=emp_h, json={"vendor_name": "Otro nombre"})
    month = today.strftime("%Y-%m")

    await harness.set_setting("usage.included_documents", 0)
    usage = (await harness.client.get("/api/v1/usage", headers=acc_h, params={"month": month})).json()
    assert usage["documents"] == 1 and usage["extra_documents"] == 1 and Decimal(usage["extra_amount"]) == Decimal("350")
    assert Decimal(usage["cost_usd"]) == Decimal("0.002")
    assert (await harness.client.get("/api/v1/usage", headers=acc_h, params={"month": "octubre"})).status_code == 422

    accuracy = (
        await harness.client.get(f"/api/v1/companies/{company.id}/ai-accuracy", headers=acc_h, params={"month": month})
    ).json()
    assert accuracy["documents"] == 1 and accuracy["fields"]["vendor_name"] == 0.0 and accuracy["fields"]["total"] == 1.0
    assert accuracy["corrected_before_confirm"] == 1


async def test_ai_costs_by_month_model_and_company_for_the_platform_operator(harness: Harness) -> None:
    company, other, _, acc_h, emp_h, _, today = await _scenario(harness)
    operator = await harness.add_employee(company, "ops@nextbyte.com", Role.SUPERADMIN)
    ops_h = await harness.login(operator)
    await _upload_and_confirm(harness, emp_h, 5, total=Decimal("45000"), invoice_date=today)
    await _upload_and_confirm(harness, emp_h, 6, total=Decimal("46000"), invoice_date=today)
    # Una de las lecturas pasa al mes anterior y con otro modelo.
    previous = (today.replace(day=1) - timedelta(days=1)).strftime("%Y-%m")
    async with harness.container.engine.begin() as conn:
        await conn.execute(
            sa.text(
                "UPDATE extraction_attempts SET created_at = date_trunc('month', now()) - interval '1 day', model_used = 'other-model', "
                "prompt_tokens = 1000, completion_tokens = 200, cached_tokens = 300 "
                "WHERE id = (SELECT id FROM extraction_attempts ORDER BY created_at LIMIT 1)"
            )
        )

    assert (await harness.client.get("/api/v1/ai-costs", headers=acc_h)).status_code == 403
    report = (await harness.client.get("/api/v1/ai-costs", headers=ops_h)).json()
    assert len(report["months"]) == 6 and report["months"][0]["month"] == today.strftime("%Y-%m")
    # La segunda factura escaló al modelo fuerte: tres llamadas en total.
    assert report["totals"]["calls"] == 3 and Decimal(report["totals"]["cost_usd"]) == Decimal("0.006")
    current, before = report["months"][0], next(m for m in report["months"] if m["month"] == previous)
    assert current["totals"]["calls"] == 2
    assert sorted(m["model"] for m in current["models"]) == ["primary-model", "strong-model"]
    assert current["models"][0]["companies"][0]["company_name"] == company.name
    assert before["models"][0]["model"] == "other-model"
    assert before["totals"]["prompt_tokens"] == 1000 and before["totals"]["cached_tokens"] == 300
    assert all(m["totals"]["calls"] == 0 for m in report["months"] if m["month"] not in (current["month"], previous))

    only_previous = await harness.client.get("/api/v1/ai-costs", headers=ops_h, params={"from": previous, "to": previous})
    assert [m["month"] for m in only_previous.json()["months"]] == [previous]
    wrong = await harness.client.get("/api/v1/ai-costs", headers=ops_h, params={"from": "2026-10", "to": "2026-09"})
    assert wrong.status_code == 422 and wrong.json()["error"]["code"] == "invalid_month_range"

    excel = await harness.client.get("/api/v1/ai-costs", headers=ops_h, params={"format": "xlsx"})
    sheet = load_workbook(io.BytesIO(excel.content)).active
    assert sheet is not None
    values = [c.value for row in sheet.iter_rows() for c in row]
    assert "other-model" in values and "Tokens en caché" in values
    assert other.name not in values
