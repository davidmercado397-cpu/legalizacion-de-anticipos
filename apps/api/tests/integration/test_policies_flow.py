from __future__ import annotations

import io
from decimal import Decimal

from openpyxl import load_workbook

from app.domain.enums import DocumentType, Role
from app.domain.values import InvoiceLine
from tests.conftest import Harness, wa_message
from tests.fakes import invoice, make_jpeg

PHONE = "573001112233"
BSUID = "CO.1502576394655843"

HOTEL = dict(
    document_type=DocumentType.ELECTRONIC_INVOICE,
    vendor_name="Hotel Dann",
    invoice_number="FE-4521",
    category="lodging",
    subtotal=Decimal("564000"),
    tax_amount=Decimal("107160"),
    total=Decimal("671160"),
    lines=(
        InvoiceLine("Habitación 3 noches", Decimal("540000")),
        InvoiceLine("Minibar cerveza x2", Decimal("24000"), concept_code="minibar", concept_confidence=0.92),
    ),
)


async def _scenario(harness: Harness):  # type: ignore[no-untyped-def]
    company = await harness.add_company()
    accountant = await harness.add_employee(company, "conta@forsa.com", Role.ACCOUNTANT)
    approver = await harness.add_employee(company, "ana@forsa.com", Role.APPROVER)
    employee = await harness.add_employee(company, "juan@forsa.com", phone=PHONE, approver=approver, wa_user_id=BSUID)
    emp_h, appr_h, acc_h = await harness.login(employee), await harness.login(approver), await harness.login(accountant)
    advance = await harness.client.post(
        "/api/v1/advances",
        headers=emp_h,
        json={
            "company_id": str(company.id),
            "kind": "cash",
            "currency": "COP",
            "amount": "800000",
            "purpose": "Obra Medellín",
            "destination": "Medellín",
            "start_date": "2026-10-12",
            "end_date": "2026-10-15",
        },
    )
    await harness.client.post(f"/api/v1/advances/{advance.json()['id']}/approve", headers=appr_h)
    await harness.drain()
    harness.wa.sent.clear()
    return company, employee, emp_h, appr_h, acc_h, advance.json()["id"]


async def _upload(harness: Harness, headers: dict[str, str], seed: int = 1) -> dict:  # type: ignore[type-arg]
    response = await harness.client.post(
        "/api/v1/expenses", headers=headers, files={"file": ("f.jpg", make_jpeg(seed=seed), "image/jpeg")}
    )
    await harness.drain()
    return (await harness.client.get(f"/api/v1/expenses/{response.json()['id']}", headers=headers)).json()


async def test_minibar_needs_justification_then_approver_reviews(harness: Harness) -> None:
    company, employee, emp_h, appr_h, acc_h, advance_id = await _scenario(harness)
    harness.extractor.respond = lambda _r, _n: [invoice(**HOTEL)]
    expense = await _upload(harness, emp_h)
    assert {f["code"] for f in expense["policy_findings"]} == {"blocked_concept"}
    assert expense["suggested_non_reimbursable"] == "28560.00"
    assert expense["review_status"] == "pending" and expense["deductibility"] == "deductible"
    lines = (await harness.client.get(f"/api/v1/expenses/{expense['id']}/lines", headers=emp_h)).json()["data"]
    assert [line["final_concept_code"] for line in lines] == [None, "minibar"]
    # La IA recibió los conceptos de la compañía en el contexto apto para caché.
    assert {c.code for c in harness.extractor.calls[0].context.blocked_concepts} == {"alcohol", "minibar", "tobacco"}

    missing = await harness.client.post(f"/api/v1/expenses/{expense['id']}/confirm", headers=emp_h, json={})
    assert missing.status_code == 409 and missing.json()["error"]["code"] == "justification_required"
    confirmed = await harness.client.post(
        f"/api/v1/expenses/{expense['id']}/confirm", headers=emp_h, json={"justification": "El minibar fue un error, lo asumo"}
    )
    assert confirmed.status_code == 200

    assert (await harness.client.post(f"/api/v1/advances/{advance_id}/submit", headers=emp_h)).status_code == 200
    blocked_close = await harness.client.post(f"/api/v1/advances/{advance_id}/close", headers=appr_h)
    assert blocked_close.status_code == 409  # primero hay que resolver la revisión pendiente

    own_review = await harness.client.post(
        f"/api/v1/expenses/{expense['id']}/review", headers=emp_h, json={"decision": "accept_full", "note": "yo"}
    )
    assert own_review.status_code == 403
    no_note = await harness.client.post(
        f"/api/v1/expenses/{expense['id']}/review", headers=appr_h, json={"decision": "accept_full"}
    )
    assert no_note.status_code == 409
    reviewed = await harness.client.post(
        f"/api/v1/expenses/{expense['id']}/review", headers=appr_h, json={"decision": "accept_without_blocked"}
    )
    assert reviewed.json()["review_status"] == "accepted_partial" and reviewed.json()["non_reimbursable_amount"] == "28560.00"

    detail = (await harness.client.get(f"/api/v1/advances/{advance_id}", headers=emp_h)).json()
    assert Decimal(detail["balance"]["applied"]) == Decimal("642600")  # 671.160 - 28.560
    assert (await harness.client.post(f"/api/v1/advances/{advance_id}/close", headers=appr_h)).json()["status"] == "closed"

    export = await harness.client.get(f"/api/v1/advances/{advance_id}/legalization.xlsx", headers=acc_h)
    assert export.status_code == 200 and "attachment" in export.headers["content-disposition"]
    sheet = load_workbook(io.BytesIO(export.content)).active
    values = [[c for c in row] for row in sheet.iter_rows(values_only=True)]
    header_index = next(i for i, row in enumerate(values) if row and row[0] == "Fecha")
    row = dict(zip(values[header_index], values[header_index + 1], strict=True))
    assert row["Proveedor"] == "Hotel Dann" and row["Total"] == 671160
    assert row["No reembolsable"] == 28560 and row["Neto legalizado"] == 642600
    assert row["Categoría"] == "Hospedaje"


async def test_whatsapp_detail_explains_the_finding_and_asks_for_a_justification(harness: Harness) -> None:
    await _scenario(harness)
    harness.extractor.respond = lambda _r, _n: [invoice(**HOTEL)]
    harness.wa.media["M1"] = (make_jpeg(), "image/jpeg")
    await harness.webhook(
        wa_message("wamid.H1", phone=PHONE, user_id=BSUID, type="image", image={"id": "M1", "mime_type": "image/jpeg"})
    )
    await harness.drain()
    detail = harness.wa.of_kind("buttons")[-1]
    assert "minibar" in detail.text and "justificación" in detail.text
    assert [b.title for b in detail.extra] == ["✍️ Justificar", "✏️ Corregir"]
    assert harness.wa.of_kind("reaction")[-1].text == "⚠️"


async def test_duplicate_across_employees_is_flagged_without_revealing_who(harness: Harness) -> None:
    company, employee, emp_h, appr_h, acc_h, _ = await _scenario(harness)
    colleague = await harness.add_employee(company, "pedro@forsa.com")
    col_h = await harness.login(colleague)
    first = await _upload(harness, emp_h, seed=1)
    assert first["policy_findings"] == []
    second = await _upload(harness, col_h, seed=2)
    duplicate = next(f for f in second["policy_findings"] if f["code"] == "duplicate_invoice")
    assert duplicate["detail"] == {} and str(employee.id) not in str(second)


async def test_accounting_configures_rules_and_concepts(harness: Harness) -> None:
    company, employee, emp_h, appr_h, acc_h, _ = await _scenario(harness)
    rules_url = f"/api/v1/companies/{company.id}/policy-rules"
    defaults = {r["rule_type"]: r for r in (await harness.client.get(rules_url, headers=emp_h)).json()["data"]}
    assert defaults["blocked_concept"]["action"] == "require_justification"

    forbidden = await harness.client.put(rules_url, headers=emp_h, json=[{"rule_type": "duplicate", "action": "block"}])
    assert forbidden.status_code == 403
    bad = await harness.client.put(
        rules_url, headers=acc_h, json=[{"rule_type": "duplicate", "action": "block", "params": {"x": 1}}]
    )
    assert bad.status_code == 422
    updated = await harness.client.put(rules_url, headers=acc_h, json=[{"rule_type": "blocked_concept", "action": "block"}])
    assert {r["rule_type"]: r for r in updated.json()["data"]}["blocked_concept"]["customized"] is True

    concepts_url = f"/api/v1/companies/{company.id}/blocked-concepts"
    created = await harness.client.post(
        concepts_url, headers=acc_h, json={"code": "spa", "name": "Spa y masajes", "examples": ["masaje", "spa"]}
    )
    assert created.status_code == 201
    assert (await harness.client.post(concepts_url, headers=acc_h, json={"code": "spa", "name": "Otro"})).status_code == 409
    assert len((await harness.client.get(concepts_url, headers=emp_h)).json()["data"]) == 4

    harness.extractor.respond = lambda _r, _n: [invoice(**HOTEL)]
    expense = await _upload(harness, emp_h)
    assert "spa" in {c.code for c in harness.extractor.calls[-1].context.blocked_concepts}
    blocked = await harness.client.post(f"/api/v1/expenses/{expense['id']}/confirm", headers=emp_h, json={"justification": "x"})
    assert blocked.status_code == 409 and blocked.json()["error"]["code"] == "policy_blocked"
    review = await harness.client.post(
        f"/api/v1/expenses/{expense['id']}/review", headers=appr_h, json={"decision": "accept_full", "note": "Cliente invitó"}
    )
    assert review.status_code == 200
    assert (
        await harness.client.post(f"/api/v1/expenses/{expense['id']}/confirm", headers=emp_h, json={"justification": "x"})
    ).status_code == 200


async def test_reviewer_corrects_a_line_the_ai_missed(harness: Harness) -> None:
    company, employee, emp_h, appr_h, acc_h, _ = await _scenario(harness)
    harness.extractor.respond = lambda _r, _n: [
        invoice(
            total=Decimal("60000"),
            subtotal=None,
            tax_amount=None,
            lines=(InvoiceLine("Almuerzo ejecutivo", Decimal("45000")), InvoiceLine("Club Colombia Roja", Decimal("15000"))),
        )
    ]
    expense = await _upload(harness, emp_h)
    assert expense["policy_findings"] == []
    lines = (await harness.client.get(f"/api/v1/expenses/{expense['id']}/lines", headers=appr_h)).json()["data"]
    unknown = await harness.client.patch(
        f"/api/v1/expenses/{expense['id']}/lines/{lines[1]['id']}", headers=appr_h, json={"concept_code": "caviar"}
    )
    assert unknown.status_code == 422
    fixed = await harness.client.patch(
        f"/api/v1/expenses/{expense['id']}/lines/{lines[1]['id']}", headers=appr_h, json={"concept_code": "alcohol"}
    )
    assert fixed.status_code == 200 and fixed.json()["corrected_by"]
    after = (await harness.client.get(f"/api/v1/expenses/{expense['id']}", headers=emp_h)).json()
    assert {f["code"] for f in after["policy_findings"]} == {"blocked_concept"}
    assert after["suggested_non_reimbursable"] == "15000.00"


async def test_cash_receipt_deductibility_and_override(harness: Harness) -> None:
    company, employee, emp_h, appr_h, acc_h, advance_id = await _scenario(harness)
    harness.extractor.respond = lambda _r, _n: [
        invoice(
            document_type=DocumentType.CASH_RECEIPT, vendor_tax_id=None, subtotal=None, tax_amount=None, lines=(), itemized=False
        )
    ]
    expense = await _upload(harness, emp_h)
    assert expense["deductibility"] == "requires_support_document" and expense["policy_findings"] == []
    url = f"/api/v1/expenses/{expense['id']}/deductibility"
    assert (await harness.client.post(url, headers=emp_h, json={"deductibility": "deductible"})).status_code == 403
    changed = await harness.client.post(url, headers=acc_h, json={"deductibility": "non_deductible"})
    assert changed.json()["deductibility"] == "non_deductible"

    pending_submit = await harness.client.post(f"/api/v1/advances/{advance_id}/submit", headers=emp_h)
    assert pending_submit.status_code == 409  # el gasto extraído se debe confirmar o descartar primero

    admin = await harness.add_employee(company, "admin@forsa.com", Role.ADMIN)
    admin_h = await harness.login(admin)
    invalid = await harness.client.patch(
        f"/api/v1/companies/{company.id}", headers=admin_h, json={"deductibility": {"pos_ticket": "maybe"}}
    )
    assert invalid.status_code == 422


async def test_other_employees_cannot_export(harness: Harness) -> None:
    company, employee, emp_h, appr_h, acc_h, advance_id = await _scenario(harness)
    other = await harness.add_employee(company, "pedro@forsa.com")
    response = await harness.client.get(f"/api/v1/advances/{advance_id}/legalization.xlsx", headers=await harness.login(other))
    assert response.status_code == 404
