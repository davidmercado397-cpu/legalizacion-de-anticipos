"""Presupuestos de viaje: zonas y tarifas (contabilidad) y cálculo del anticipo de un viaje (empleado)."""

from __future__ import annotations

import io
import uuid
import zipfile
from typing import Annotated, Any

from fastapi import APIRouter, File, UploadFile, status
from openpyxl import load_workbook

from app.adapters.inbound.http.deps import ContainerDep, CurrentEmployee
from app.adapters.inbound.http.schemas import (
    BudgetLineOut,
    BudgetOut,
    ImportOut,
    ListOut,
    TravelRateIn,
    TravelRateOut,
    TravelZoneIn,
    TravelZoneOut,
    TravelZonePatch,
    TripIn,
)
from app.application.budgets import RateData, TripQuery
from app.domain.errors import ValidationError

router = APIRouter(tags=["travel budgets"])
_COMPANY = "/companies/{company_id}"


@router.get(f"{_COMPANY}/travel-zones", response_model=ListOut[TravelZoneOut])
async def list_zones(company_id: uuid.UUID, actor: CurrentEmployee, container: ContainerDep) -> ListOut[TravelZoneOut]:
    zones = await container.budgets.zones(actor, company_id)
    return ListOut[TravelZoneOut](data=[TravelZoneOut.model_validate(z) for z in zones])


@router.post(f"{_COMPANY}/travel-zones", response_model=TravelZoneOut, status_code=status.HTTP_201_CREATED)
async def add_zone(company_id: uuid.UUID, body: TravelZoneIn, actor: CurrentEmployee, container: ContainerDep) -> TravelZoneOut:
    return TravelZoneOut.model_validate(await container.budgets.add_zone(actor, company_id, body.name, body.places))


@router.patch(f"{_COMPANY}/travel-zones/{{zone_id}}", response_model=TravelZoneOut)
async def update_zone(
    company_id: uuid.UUID, zone_id: uuid.UUID, body: TravelZonePatch, actor: CurrentEmployee, container: ContainerDep
) -> TravelZoneOut:
    zone = await container.budgets.update_zone(actor, company_id, zone_id, body.model_dump(exclude_unset=True))
    return TravelZoneOut.model_validate(zone)


@router.get(f"{_COMPANY}/travel-rates", response_model=ListOut[TravelRateOut], summary="Travel rates (accounting only)")
async def list_rates(
    company_id: uuid.UUID, actor: CurrentEmployee, container: ContainerDep, zone_id: uuid.UUID | None = None
) -> ListOut[TravelRateOut]:
    rates = await container.budgets.rates(actor, company_id, zone_id)
    return ListOut[TravelRateOut](data=[TravelRateOut.model_validate(r) for r in rates])


@router.post(f"{_COMPANY}/travel-rates", response_model=TravelRateOut, status_code=status.HTTP_201_CREATED)
async def add_rate(company_id: uuid.UUID, body: TravelRateIn, actor: CurrentEmployee, container: ContainerDep) -> TravelRateOut:
    rate = await container.budgets.add_rate(actor, company_id, RateData(**body.model_dump()))
    return TravelRateOut.model_validate(rate)


@router.delete(f"{_COMPANY}/travel-rates/{{rate_id}}", status_code=status.HTTP_204_NO_CONTENT, summary="Deactivate a rate")
async def deactivate_rate(company_id: uuid.UUID, rate_id: uuid.UUID, actor: CurrentEmployee, container: ContainerDep) -> None:
    await container.budgets.deactivate_rate(actor, company_id, rate_id)


@router.post(
    f"{_COMPANY}/travel-rates/import",
    response_model=ImportOut,
    summary="Load rates from Excel (columns: zona, tipo_personal, categoria, valor, moneda, unidad, desde, hasta)",
)
async def import_rates(
    company_id: uuid.UUID,
    actor: CurrentEmployee,
    container: ContainerDep,
    file: Annotated[UploadFile, File(description=".xlsx")],
) -> ImportOut:
    content = await file.read(container.settings.MAX_UPLOAD_BYTES + 1)
    if len(content) > container.settings.MAX_UPLOAD_BYTES:
        raise ValidationError(
            "The file is too large", code="file_too_large", max_mb=container.settings.MAX_UPLOAD_BYTES // (1024 * 1024)
        )
    return ImportOut(imported=await container.budgets.import_rates(actor, company_id, _rows(content)))


@router.post("/advances/budget-preview", response_model=BudgetOut, summary="Suggested advance for a trip")
async def budget_preview(body: TripIn, actor: CurrentEmployee, container: ContainerDep) -> BudgetOut:
    legs = tuple(leg.to_domain() for leg in body.legs)
    query = TripQuery(body.company_id, body.destination, body.start_date, body.end_date, body.travel_zone_id, legs)
    budget = await container.budgets.preview(actor, query)
    lines = [BudgetLineOut(**line.as_dict()) for line in budget.lines]
    return BudgetOut(zone_id=budget.zone_id, currency=budget.currency, total=budget.total, lines=lines)


def _rows(content: bytes) -> list[dict[str, Any]]:
    """Lee la primera hoja: la fila 1 son los nombres de columna (sin importar mayúsculas)."""
    try:
        sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
    except (zipfile.BadZipFile, KeyError, OSError, ValueError, IndexError) as exc:
        raise ValidationError("Invalid spreadsheet", code="invalid_spreadsheet") from exc
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return []
    header = [str(cell or "").strip().lower() for cell in rows[0]]
    return [dict(zip(header, row, strict=False)) for row in rows[1:] if any(cell not in (None, "") for cell in row)]
