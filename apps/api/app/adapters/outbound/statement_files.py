"""Lectura de extractos bancarios en CSV o Excel como filas de texto (el mapeo de columnas lo hace el dominio)."""

from __future__ import annotations

import csv
import io
from datetime import date, datetime
from decimal import Decimal

from openpyxl import load_workbook

from app.domain.errors import ValidationError

_MAX_ROWS = 5000


class SpreadsheetTableReader:
    def read(self, content: bytes, file_name: str) -> list[list[str]]:
        name = file_name.lower()
        if name.endswith((".xlsx", ".xlsm")) or content[:2] == b"PK":
            return _excel(content)
        if name.endswith((".csv", ".txt")):
            return _csv(content)
        raise ValidationError("Unsupported statement format", code="statement_unsupported_format")


def _excel(content: bytes) -> list[list[str]]:
    try:
        book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - un archivo dañado es un error del usuario, no del servidor
        raise ValidationError("Unreadable Excel file", code="statement_unreadable") from exc
    sheet = book.worksheets[0]
    rows = [[_cell(value) for value in row] for row in sheet.iter_rows(values_only=True, max_row=_MAX_ROWS)]
    book.close()
    return rows


def _cell(value: object) -> str:
    """Las fechas de Excel salen en ISO y los números con punto decimal: el perfil los reconoce igual."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float):
        return format(Decimal(str(value)).normalize(), "f")
    return str(value).strip()


def _csv(content: bytes) -> list[list[str]]:
    for encoding in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    # El separador es el que más se repite en las primeras líneas (los bancos ponen títulos sueltos arriba).
    lines = text.splitlines()[:30]
    delimiter = max(";,\t|", key=lambda d: sum(line.count(d) for line in lines))
    return [[c.strip() for c in row] for row in list(csv.reader(io.StringIO(text), delimiter=delimiter))[:_MAX_ROWS]]
